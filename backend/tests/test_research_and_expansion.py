"""Title I parsing, pilot scoring, translation fallback, and configuration-only jurisdiction expansion.
Spreadsheets are synthetic, built in-test with the real column layout (no network)."""
import io
import shutil

import yaml
from openpyxl import Workbook
from sqlalchemy import select

from app.config import REPO_ROOT
from app.ingestion.base import Fetcher, SourceConfig
from app.ingestion.pipeline import run_source
from app.models import Jurisdiction, Location, SchoolDivision, Translation
from app.research.pilot_selection import score_divisions
from app.research.title_i import division_summary, parse_free_reduced, parse_title_i
from app.services.registry import active_scope, sync_jurisdictions
from app.services.translation import pick_text, translate_pending


def _title_i_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "DATA"
    ws.append(["DIVISION NAME", "DIVISION NUMBER", "SCHOOL NAME", "SCHOOL NUMBER", "GRADE SPAN SERVED", "SCHOOL TYPE", "FOCUS AREA"])
    ws.append(["Synthetic A", 1, "School One ", 10, "PK-05", "Elementary", "Title I - Schoolwide"])
    ws.append(["Synthetic A", 1, "School Two", 11, "PK-05", "Elementary", "Title I - Targeted Assistance"])
    ws.append(["Synthetic B", 2, "School Nine", 90, "06-08", "Middle", "Title I - Schoolwide"])
    wb.create_sheet("SQL")
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _fr_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    for _ in range(6):
        ws.append(["title row"])
    ws.append(["001-Synthetic A"])
    ws.append(["Div. No.", "School No.", "School Name", "CEP*", "School Type", "Low", "High", "SNP Memb", "FREE Elig", "FREE %", "RED Elig", "RED %", "TOTAL F/R Elig", "TOTAL F/R %"])
    for sch, cep, mem, fr in (("0010", "X", 100, 80), ("0011", "", 100, 20), ("0012", "", 200, 20), ("0013", "", 100, 10)):
        ws.append(["001", sch, "S", cep, "Elem", "PK", "05", mem, 0, 0, 0, 0, fr, 0])
    ws.append(["", "", "TOTALS"])
    ws.append(["002", "0090", "S9", "", "Mid", "06", "08", 50, 0, 0, 0, 0, 25, 0])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_title_i_and_free_reduced_parsing_and_summary():
    ti = parse_title_i(_title_i_xlsx())
    fr = parse_free_reduced(_fr_xlsx())
    assert list(ti["div"]) == [1, 1, 2] and ti.iloc[0]["school_name"] == "School One"
    assert len(fr) == 5 and fr.iloc[0]["cep"]
    acs = {"001": {"households": 100, "snap_households": 10, "pov_universe": 200, "below_poverty": 20, "below_185pct": 50}}
    divisions = {1: {"name": "Synthetic A", "fips": ["001"], "jurisdictions": ["a"]}, 2: {"name": "Synthetic B", "fips": ["999"], "jurisdictions": ["b"]}}
    rows = {r["division_number"]: r for r in division_summary(ti, fr, acs, divisions, {"title_i": {"school_year": "2025-2026"}})}
    a = rows[1]
    assert (a["total_schools"], a["title_i_schools"], a["title_i_share"]) == (4, 2, 0.5)
    assert a["fr_eligible_pct"] == 26.0 and a["cep_schools"] == 1
    assert a["acs_snap_household_pct"] == 10.0 and a["acs_below_185pct_fpl_pct"] == 25.0
    assert any("Community Eligibility" in w for w in a["warnings"])
    assert "acs_snap_household_pct" not in rows[2] and any("unavailable" in w for w in rows[2]["warnings"])  # never imputed


def test_title_i_missing_columns_fails_loudly():
    wb = Workbook()
    wb.active.title = "DATA"
    wb.active.append(["DIVISION NAME", "SCHOOL NAME"])
    b = io.BytesIO()
    wb.save(b)
    import pytest

    with pytest.raises(ValueError):
        parse_title_i(b.getvalue())


def test_pilot_scoring_normalizes_and_excludes_missing():
    rows = [
        {"division_number": 1, "division_name": "A", "jurisdictions": ["a"], "title_i_schools": 10, "title_i_share": 0.5},
        {"division_number": 2, "division_name": "B", "jurisdictions": ["b"], "title_i_schools": 0, "title_i_share": 0.1},
        {"division_number": 3, "division_name": "C", "jurisdictions": ["c"], "title_i_schools": 5, "title_i_share": None},
    ]
    cfg = {"candidates": [1, 2, 3], "factors": {"n": {"weight": 1, "field": "title_i_schools"}, "s": {"weight": 1, "field": "title_i_share"}}}
    out = {r["division_number"]: r for r in score_divisions(rows, cfg)}
    assert out[1]["score"] == 1.0 and out[2]["score"] == 0.0
    assert out[3]["score"] == 0.5 and out[3]["missing_factors"] == ["s"] and out[3]["weight_coverage"] == 0.5
    assert out[1]["rank"] == 1


def test_translation_pipeline_and_fallback(db, imported):
    calls = []

    def fake(text, src, tgt):
        calls.append(text)
        return f"[{tgt}] {text}"

    res = translate_pending(db, "es", translator=fake)
    assert res["translated"] > 0 and calls
    t = db.scalar(select(Translation).where(Translation.lang == "es"))
    assert t.is_machine and not t.reviewed
    again = translate_pending(db, "es", translator=fake)
    assert again["translated"] == 0  # unchanged originals are not re-sent
    assert pick_text("Hello", "en", t, "es")["machine"] is True
    assert pick_text("Hello", "en", None, "es") == {"text": "Hello", "lang": "en", "machine": False, "is_original": True}
    assert pick_text("Hola", "es", t, "es")["is_original"]
    assert translate_pending(db, "es")["status"] == "skipped"  # provider 'none' by default: nothing sent anywhere


def test_new_jurisdiction_added_by_configuration_only(db, tmp_path, monkeypatch):
    """Expansion: copy config, add a jurisdiction YAML + enable it — no code changes, no DB reinit."""
    cfg_dir = tmp_path / "config"
    shutil.copytree(REPO_ROOT / "config", cfg_dir)
    (cfg_dir / "jurisdictions" / "test-town.yaml").write_text(yaml.safe_dump({
        "id": "test-town", "name": "Test Town (synthetic)", "kind": "town", "county_fips": None, "status": "expansion",
        "boundary": {"source": "none"}, "school_divisions": [{"vdoe_number": 999, "name": "Synthetic Division", "slug": "synthetic"}],
    }))
    before = db.scalar(select(Location.id))
    sync_jurisdictions(db, cfg_dir)
    assert db.get(Jurisdiction, "test-town").status == "expansion"
    assert db.get(SchoolDivision, 999).jurisdictions[0].id == "test-town"
    assert "test-town" in active_scope(cfg_dir)
    f = tmp_path / "t.json"
    f.write_text('{"records":[{"external_id":"t1","organization_name":"[TEST] Town Pantry","address1":"1 Main","zip":"22000","lat":38.5,"lon":-77.6,"jurisdiction_id":"test-town"}]}')
    src = SourceConfig(id="syn-town", name="t", url="file://t", adapter="structured_json", jurisdictions=["test-town"], is_synthetic=True)
    run = run_source(db, src, Fetcher(fixtures={"file://t": f}), scope=active_scope(cfg_dir))
    assert run.counts["inserted"] == 1
    assert db.scalar(select(Location).where(Location.jurisdiction_id == "test-town")) is not None
    assert before is None or db.get(Location, before) is not None  # existing data untouched
    # re-sync is idempotent
    sync_jurisdictions(db, cfg_dir)
    assert db.scalar(select(Jurisdiction).where(Jurisdiction.id == "test-town")).name == "Test Town (synthetic)"


def test_pilot_config_changes_scope_without_code(tmp_path):
    cfg_dir = tmp_path / "config"
    shutil.copytree(REPO_ROOT / "config", cfg_dir)
    (cfg_dir / "pilot.yaml").write_text("pilot_jurisdictions: [arlington-county]\n")
    scope = active_scope(cfg_dir)
    assert "arlington-county" in scope and "prince-william-county" not in scope  # pilot list replaced; expansion kept

